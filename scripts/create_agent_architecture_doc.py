from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "산출물 1 에이전트 아키텍처 설계서.docx"
FLOW_IMAGE = Path(r"C:\Users\Playdata\Downloads\diagram (1).jpg")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        header.cells[index].width = Cm(widths[index])
        set_cell_shading(header.cells[index], "1F4E78")
        set_cell_text(header.cells[index], label, bold=True, color="FFFFFF")
        header.cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Cm(widths[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F4F8FC")
            set_cell_text(cells[index], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Malgun Gothic"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(7)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def build_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("LawPath 에이전트 아키텍처 설계서")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("법률 정보 검색과 근거 기반 응답 생성을 위한 실행 구조")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(68, 68, 68)

    doc.add_paragraph(
        "이 문서는 LawPath 프로젝트의 현재 구현을 기준으로, 사용자 질문이 입력 판단, "
        "도구 선택, 법률 근거 검색, 응답 생성, 저장 및 화면 표시로 이어지는 에이전트 실행 구조를 정의한다. "
        "핵심 원칙은 검색된 Evidence만으로 응답을 구성하고, 정보가 부족하면 검색 전에 보완 질문으로 종료하는 것이다."
    )

    doc.add_heading("1 서비스 개요", level=1)
    add_table(doc, ["항목", "내용"], [
        ["서비스명", "LawPath"],
        ["목적", "주거, 근로 임금, 소비자 분쟁 질문에 대해 법령, 판례, 상담사례를 검색하고 근거 기반 안내를 제공"],
        ["주요 사용자", "생활 법률 정보를 확인하려는 비회원 및 회원 사용자"],
        ["입력 채널", "Streamlit 웹 화면의 일반 분석 요청, 후속 질문, 법률 용어 대화"],
        ["연계 구성", "FastAPI Backend, Legal MCP, PostgreSQL pgvector, LLM Provider, SSE 실행 이벤트"],
        ["안전 원칙", "Evidence 밖의 법령, 판례, 사실 또는 결론을 추가하지 않고 법률 자문이나 결과 보장을 하지 않음"],
    ], [3.1, 12.3])

    doc.add_heading("2 StateGraph 노드 설계", level=1)
    doc.add_paragraph(
        "구현은 단일 LangGraph 객체가 아니라 AgentRuntime과 서비스 계층으로 구성되어 있다. "
        "아래 표는 현재 실행 단계를 StateGraph 관점의 노드로 정리한 것이다."
    )
    add_table(doc, ["노드", "역할", "주요 입력", "주요 출력과 다음 단계"], [
        ["요청 수신", "질문, 카테고리, 사용자 식별자와 저장 선택을 수신하고 idempotency key를 확인", "session id, category, question, save_selected", "run 생성 또는 기존 run 재사용 -> 입력 판단"],
        ["입력 판단", "LLM 구조화 출력으로 검색 준비 상태를 판단", "category, question, 이전 대화", "sufficient, proceed_with_caution 또는 needs_clarification"],
        ["보완 응답", "검색하지 않고 부족한 정보와 최대 3개 보완 질문을 반환", "IntakeResult", "status=stopped, termination_reason=needs_clarification -> 종료"],
        ["프로필 선택", "카테고리별 AgentProfile과 허용 도구를 결정", "housing, labor, consumer", "profile, allowed_tools -> 도구 선택"],
        ["도구 선택", "카테고리에 맞는 정해진 검색 순서를 적용", "profile, question", "labor housing은 판례와 통합 법률자료, consumer는 법령 상담사례 판례"],
        ["도구 호출", "Legal MCP 검색 도구를 순차 호출하고 SSE trace를 발행", "query, category, top_k=3", "Evidence 목록, tool_completed event -> 근거 정규화"],
        ["근거 정규화", "문서 식별자로 중복을 제거하고 근거 수와 종료 사유를 계산", "raw Evidence", "no_results, insufficient_evidence 또는 model_finished -> 답변 생성"],
        ["답변 생성", "LLM 구조화 응답을 검증하고 실패 시 템플릿 AnswerAgent로 복귀", "question, 제한된 Evidence, 대화 컨텍스트", "AnswerDraft, used_evidence_ids -> 응답 조립"],
        ["응답 저장 및 발행", "명시적 저장 선택 시 회원 대화를 저장하고 terminal SSE 이벤트를 발행", "response, actor, save_selected", "completed 또는 stopped 결과 -> UI"],
    ], [2.5, 4.0, 4.0, 4.9])

    doc.add_page_break()
    doc.add_heading("3 상태 흐름 예시", level=1)
    doc.add_paragraph(
        "제공된 흐름도는 LawPath의 입력 판단과 검색 실행 흐름에 대응한다. "
        "현재 구현에서 일정 생성 단계는 사용하지 않으며, 후보 제시 이후의 동작은 근거 기반 답변과 후속 질문 표시로 대체된다."
    )
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(FLOW_IMAGE), width=Cm(9.2))
    add_caption(doc, "그림 1 요청 분석부터 도구 호출과 재시도까지의 상태 흐름")
    add_bullet(doc, "요청 분석은 IntakeAgent의 구조화 입력 판단에 해당한다.")
    add_bullet(doc, "필수값 충분 여부가 아니오이면 needs_clarification 응답과 보완 질문을 반환하고 검색 도구를 호출하지 않는다.")
    add_bullet(doc, "도구 선택 호출은 LegalAgentRuntime의 profile 기반 검색 순서와 MCP 호출에 해당한다.")
    add_bullet(doc, "결과 유효성은 MCP 성공 여부, Evidence 형식, 문서 중복 제거 및 근거 수로 판단한다.")
    add_bullet(doc, "성찰 수정 재시도는 현재 구현에서는 입력 판단의 형식 오류 1회 재시도와 도구 오류의 실패 응답 처리로 제한된다.")

    doc.add_heading("4 주요 분기 조건과 예외 폴백", level=1)
    add_table(doc, ["조건", "처리", "사용자 또는 시스템 결과"], [
        ["입력 판단이 needs_clarification", "검색과 MCP 호출을 건너뛰고 보완 응답 생성", "stopped 상태와 최대 3개의 후속 질문"],
        ["입력 판단이 sufficient 또는 proceed_with_caution", "검색 실행을 허용", "입력 판단 메시지와 주의사항을 유지한 채 다음 노드로 진행"],
        ["LLM 입력 판단의 계약 오류", "구조화 출력 재요청은 최대 1회, 전체 시간 제한 내에서만 수행", "실패 시 IntakeAssessmentError로 실행 실패"],
        ["허용되지 않은 도구 또는 3회 초과 도구 호출", "정책 검증에서 차단", "실행 오류로 종료"],
        ["MCP 도구 실패 또는 잘못된 data 형식", "도구 오류를 RuntimeError로 전파", "SSE run.failed 및 일반화된 오류 메시지"],
        ["검색 결과 0건", "답변은 근거 없는 내용을 만들지 않음", "termination_reason=no_results"],
        ["근거가 3건 미만", "반환 가능한 Evidence는 유지하되 주의사항 추가", "termination_reason=insufficient_evidence"],
        ["답변 LLM 실패 또는 인용 id 검증 실패", "AnswerAgent 템플릿 답변으로 폴백", "응답은 유지하고 llm_calls=0"],
        ["저장 미선택 또는 비회원", "영구 저장을 수행하지 않음", "비회원 분석은 임시 이력 저장을 시도"],
    ], [5.0, 6.0, 4.4])

    doc.add_heading("5 공유 상태 객체", level=1)
    add_table(doc, ["필드", "타입", "사용 노드", "설명"], [
        ["request_id", "string", "요청 수신, Runtime, 응답", "질문 실행과 결과를 식별"],
        ["agent_id", "string", "프로필 선택, Runtime", "housing, labor, consumer 중 하나"],
        ["question", "string", "입력 판단, 도구 호출, 답변 생성", "사용자 질문 원문"],
        ["status", "running completed failed stopped", "모든 실행 노드", "현재 실행 상태"],
        ["termination_reason", "string 또는 null", "근거 정규화, 응답 조립", "needs_clarification, no_results, insufficient_evidence, model_finished 등"],
        ["current_step", "integer", "도구 호출", "실행한 Runtime 단계 수"],
        ["llm_calls", "integer", "답변 생성", "실제 LLM 답변 생성 사용 여부"],
        ["tool_calls", "integer", "도구 호출", "MCP 도구 호출 횟수, 최대 3"],
        ["evidence_count", "integer", "근거 정규화", "중복 제거 후 Evidence 수"],
        ["trace", "list of object", "Runtime, SSE", "tool_selected, tool_completed, evidence_insufficient 이벤트"],
        ["answer", "string 또는 null", "답변 생성, UI", "최종 사용자 안내 문구"],
        ["input_assessment", "object", "입력 판단, UI", "status, message, checks를 포함하는 준비도 판단"],
    ], [3.0, 3.2, 4.1, 5.1])

    doc.add_heading("6 기억과 컨텍스트 관리", level=1)
    add_table(doc, ["구분", "현재 구현", "보존 및 축약 원칙"], [
        ["단기 상태", "Streamlit session_state에 선택 카테고리, 질문 초안, last_result, 후속 대화, SSE pending 상태를 유지", "카테고리 전환과 새 분석에서 분석 관련 화면 상태를 초기화"],
        ["실행 상태", "단일 프로세스 메모리의 agent_runs에 run, SSE events, idempotency 정보를 보관", "이벤트 id를 증가시켜 SSE 재연결 시 마지막 id 이후만 재전송"],
        ["회원 장기 기억", "명시적으로 저장한 분석과 법률용어 대화를 저장소에 보관", "소유자 검증 후에만 복원, 삭제 가능"],
        ["비회원 기억", "분석 결과 임시 저장을 시도", "영구 저장이나 회원 대화 복원은 허용하지 않음"],
        ["대화 컨텍스트", "저장 대화의 첫 사용자 메시지와 최근 3개 메시지를 선택", "전체 길이를 3000자로 제한"],
        ["LLM 입력 축약", "대화 항목은 항목당 1000자, Evidence 본문은 항목당 1500자로 제한", "오래된 원문 전체를 전달하지 않고 현재 질문과 근거 중심으로 구성"],
    ], [3.0, 7.0, 5.4])

    doc.add_heading("7 도구 명세", level=1)
    add_table(doc, ["도구", "입력", "출력", "선택 및 실패 처리"], [
        ["search_laws", "query, category, top_k", "법령 Evidence 목록", "consumer 프로필에서 우선 호출. 실패 시 SEARCH_LAWS_FAILED"],
        ["search_consultations", "query, category, top_k", "상담사례 Evidence 목록", "consumer 프로필에서 호출. 실패 시 SEARCH_CONSULTATIONS_FAILED"],
        ["search_cases", "query, category, top_k", "판례 Evidence 목록", "모든 프로필에서 사용. 실패 시 SEARCH_CASES_FAILED"],
        ["search_legal_documents", "query, category, document_types, top_k", "법령과 판례 통합 items", "housing, labor 프로필의 보조 검색. 실패 시 SEARCH_LEGAL_DOCUMENTS_FAILED"],
        ["get_law_article", "law_name, article_number", "법령 조문 상세", "허용 목록에는 포함되나 현재 Runtime의 일반 분석 흐름에서는 직접 선택하지 않음"],
        ["get_case_detail", "document_id", "판례 상세", "상세 확인용 MCP 도구. 일반 분석 Runtime의 기본 순서에는 포함되지 않음"],
    ], [3.6, 4.0, 4.0, 3.8])

    doc.add_heading("8 중복 제거와 그래프 종료 조건", level=1)
    add_bullet(doc, "Evidence 중복 제거 기준은 evidence_id, document_id, 두 값이 없을 때는 항목 표현값 순서다. 이미 본 키는 결과 목록에 다시 추가하지 않는다.")
    add_bullet(doc, "도구 호출은 프로필의 허용 도구 집합과 고정 검색 순서를 함께 만족해야 하며 최대 3회까지만 실행한다.")
    add_bullet(doc, "입력 정보가 부족하면 stopped와 needs_clarification으로 종료한다. 검색 결과가 없으면 completed와 no_results로 종료한다.")
    add_bullet(doc, "근거가 1~2건이면 completed와 insufficient_evidence로 종료하며, 3건 이상이면 model_finished로 종료한다.")
    add_bullet(doc, "도구, 입력 판단 또는 실행 서비스에서 복구할 수 없는 오류가 발생하면 failed로 전환하고 민감한 내부 오류를 사용자 응답에 노출하지 않는다.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("LawPath Agent Architecture")
    footer_run.font.name = "Malgun Gothic"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 100, 100)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
